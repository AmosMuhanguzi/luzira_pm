import csv
from datetime import date, datetime
from io import BytesIO, StringIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import (
    PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
)
from xml.sax.saxutils import escape


INMATE_EXPORT_FIELDS = (
    ('Inmate number', 'inmate_number'),
    ('Full name', 'full_name'),
    ('Aliases', 'alias_names'),
    ('Date of birth', 'date_of_birth'),
    ('Gender', 'gender'),
    ('Nationality', 'nationality'),
    ('Tribe', 'tribe'),
    ('Religion', 'religion'),
    ('National ID number', 'national_id_number'),
    ('Passport number', 'passport_number'),
    ('Next of kin name', 'next_of_kin_name'),
    ('Next of kin relationship', 'next_of_kin_relationship'),
    ('Next of kin phone', 'next_of_kin_phone'),
    ('Next of kin address', 'next_of_kin_address'),
    ('Height (cm)', 'height_cm'),
    ('Weight (kg)', 'weight_kg'),
    ('Eye color', 'eye_color'),
    ('Hair color', 'hair_color'),
    ('Distinguishing marks', 'distinguishing_marks'),
    ('Crime category', 'crime_category'),
    ('Crime description', 'crime_description'),
    ('Court case number', 'court_case_number'),
    ('Sentencing court', 'sentencing_court'),
    ('Judge name', 'judge_name'),
    ('Sentence type', 'sentence_type'),
    ('Sentence duration', 'sentence_duration'),
    ('Sentence start date', 'sentence_start_date'),
    ('Expected release date', 'expected_release_date'),
    ('Actual release date', 'actual_release_date'),
    ('Cell block', 'cell_block'),
    ('Cell number', 'cell_number'),
    ('Security classification', 'security_classification'),
    ('Status', 'status'),
    ('Current admission date', 'current_admission_date'),
    ('Total admissions', 'total_admissions'),
    ('Risk level', 'risk_level'),
    ('Violence history', 'violence_history'),
    ('Escape attempt history', 'escape_attempt_history'),
)

VISITOR_EXPORT_FIELDS = (
    ('Visitor number', 'visitor_number'),
    ('Full name', 'full_name'),
    ('National ID number', 'national_id_number'),
    ('Passport number', 'passport_number'),
    ('Gender', 'gender'),
    ('Date of birth', 'date_of_birth'),
    ('Phone number', 'phone_number'),
    ('Email', 'email'),
    ('Address', 'address'),
    ('Nationality', 'nationality'),
    ('Relationship to inmate', 'relationship_type'),
    ('Flagged', 'is_flagged'),
    ('Flag reason', 'flag_reason'),
    ('Risk rating', 'risk_rating'),
    ('Total visits', 'total_visits_made'),
    ('Anomaly flag', 'anomaly_flag'),
)


def _display_value(value):
    if value is None:
        return ''
    if isinstance(value, datetime):
        return value.isoformat(sep=' ', timespec='minutes')
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        return 'Yes' if value else 'No'
    return str(value)


def _spreadsheet_value(value):
    display = _display_value(value)
    if isinstance(value, str) and display.lstrip().startswith(('=', '+', '-', '@')):
        return "'" + display
    return display


def build_csv(records, fields):
    output = StringIO(newline='')
    writer = csv.writer(output)
    writer.writerow([label for label, _ in fields])
    for record in records:
        writer.writerow([
            _spreadsheet_value(getattr(record, attribute))
            for _, attribute in fields
        ])
    return output.getvalue().encode('utf-8-sig')


def build_excel(records, fields, title):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title[:31]
    sheet.append([label for label, _ in fields])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='343A40')
    for record in records:
        sheet.append([
            _spreadsheet_value(getattr(record, attribute))
            for _, attribute in fields
        ])
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        width = max(len(str(cell.value or '')) for cell in column)
        sheet.column_dimensions[column[0].column_letter].width = min(width + 2, 45)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def build_pdf(records, fields, title):
    output = BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=landscape(letter),
        rightMargin=0.45 * inch,
        leftMargin=0.45 * inch,
        topMargin=0.5 * inch,
        bottomMargin=0.5 * inch,
        title=title,
    )
    styles = getSampleStyleSheet()
    heading_style = ParagraphStyle(
        'RecordHeading', parent=styles['Heading2'], textColor=colors.HexColor('#212529')
    )
    cell_style = ParagraphStyle('RecordCell', parent=styles['BodyText'], fontSize=8)
    title_style = ParagraphStyle(
        'ExportTitle', parent=styles['Title'], alignment=TA_CENTER,
    )
    content = [Paragraph(escape(title), title_style), Spacer(1, 12)]
    for index, record in enumerate(records):
        record_name = _display_value(getattr(record, fields[0][1]))
        content.append(Paragraph(escape(record_name), heading_style))
        rows = [
            [
                Paragraph(f'<b>{escape(label)}</b>', cell_style),
                Paragraph(escape(_display_value(getattr(record, attribute))).replace('\n', '<br/>'), cell_style),
            ]
            for label, attribute in fields
        ]
        table = Table(rows, colWidths=[1.7 * inch, 8.2 * inch], repeatRows=0)
        table.setStyle(TableStyle([
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#F1F3F5')),
            ('GRID', (0, 0), (-1, -1), 0.35, colors.HexColor('#CED4DA')),
            ('LEFTPADDING', (0, 0), (-1, -1), 5),
            ('RIGHTPADDING', (0, 0), (-1, -1), 5),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ]))
        content.append(table)
        if index < len(records) - 1:
            content.append(PageBreak())
    if not records:
        content.append(Paragraph('No records matched the selected filters.', styles['BodyText']))
    document.build(content)
    return output.getvalue()


def build_export(records, fields, file_format, title):
    if file_format == 'csv':
        return build_csv(records, fields), 'text/csv; charset=utf-8', 'csv'
    if file_format == 'xlsx':
        return (
            build_excel(records, fields, title),
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            'xlsx',
        )
    if file_format == 'pdf':
        return build_pdf(records, fields, title), 'application/pdf', 'pdf'
    raise ValueError(f'Unsupported export format: {file_format}')
